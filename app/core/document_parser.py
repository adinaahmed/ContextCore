import fitz  # PyMuPDF
import pdfplumber
import pytesseract
from PIL import Image
import docx
from docx.document import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
import openpyxl
import csv
import json
from bs4 import BeautifulSoup
import io

from app.core.visual_understanding import describe_image

OCR_MIN_TEXT_LENGTH = 20
PDF_BOLD_FLAG = 16
PDF_HEADING_SIZE_RATIO = 1.15
PDF_HEADING_MAX_LENGTH = 120
MAX_IMAGES_PER_DOCUMENT = 10


def _iter_block_items(parent):
    if isinstance(parent, DocxDocument):
        parent_elm = parent.element.body
    else:
        parent_elm = parent._element
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def _is_docx_heading(paragraph: Paragraph) -> bool:
    if paragraph.style is None:
        return False
    style_name = (paragraph.style.name or "").lower()
    return style_name.startswith("heading") or style_name == "title"


def _format_table_as_text(rows: list) -> str:
    lines = ["TABLE:"]
    for row in rows:
        cells = [str(cell).strip() if cell is not None else "" for cell in row]
        lines.append(" | ".join(cells))
    lines.append("END TABLE")
    return "\n".join(lines)


def _format_docx_table_as_text(table: Table) -> str:
    rows = [[cell.text for cell in row.cells] for row in table.rows]
    return _format_table_as_text(rows)


def _extract_docx_images(doc: DocxDocument) -> list[bytes]:
    images = []
    for rel in doc.part.rels.values():
        if "image" in rel.reltype.lower():
            try:
                images.append(rel.target_part.blob)
            except Exception:
                continue
    return images[:MAX_IMAGES_PER_DOCUMENT]


def _extract_pptx_images(presentation: Presentation) -> list[bytes]:
    images = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    images.append(shape.image.blob)
                except Exception:
                    continue
    return images[:MAX_IMAGES_PER_DOCUMENT]


def _extract_pdf_images(fitz_doc) -> list[bytes]:
    images = []
    for page in fitz_doc:
        for img in page.get_images(full=True):
            xref = img[0]
            try:
                base_image = fitz_doc.extract_image(xref)
                images.append(base_image["image"])
            except Exception:
                continue
            if len(images) >= MAX_IMAGES_PER_DOCUMENT:
                return images
    return images


def _describe_images_as_text(images: list[bytes]) -> list[str]:
    described = []
    for image_bytes in images:
        description = describe_image(image_bytes)
        if description:
            described.append(f"[IMAGE] {description}")
    return described


def _extract_pdf_tables(pdfplumber_page) -> tuple[list[str], list[tuple]]:
    table_texts = []
    table_bboxes = []
    try:
        tables = pdfplumber_page.find_tables()
    except Exception:
        return table_texts, table_bboxes
    for table in tables:
        try:
            rows = table.extract()
            if rows and any(any(cell for cell in row) for row in rows):
                table_texts.append(_format_table_as_text(rows))
                table_bboxes.append(table.bbox)
        except Exception:
            continue
    return table_texts, table_bboxes


def _point_in_bbox(x: float, y: float, bbox: tuple) -> bool:
    x0, y0, x1, y1 = bbox
    return x0 <= x <= x1 and y0 <= y <= y1


def _detect_pdf_headings(fitz_page) -> set:
    try:
        page_dict = fitz_page.get_text("dict")
    except Exception:
        return set()
    spans = []
    for block in page_dict.get("blocks", []):
        for line in block.get("lines", []):
            for span in line.get("spans", []):
                text = span.get("text", "").strip()
                if text:
                    spans.append((text, span.get("size", 0), span.get("flags", 0)))
    if not spans:
        return set()
    sizes = sorted(s[1] for s in spans)
    body_size = sizes[len(sizes) // 2]
    headings = set()
    for text, size, flags in spans:
        is_bold = bool(flags & PDF_BOLD_FLAG)
        is_larger = size > body_size * PDF_HEADING_SIZE_RATIO
        if (is_larger or (is_bold and size >= body_size)) and len(text) < PDF_HEADING_MAX_LENGTH:
            headings.add(text)
    return headings


def _tag_if_heading(text: str, headings: set) -> str:
    stripped = text.strip()
    if stripped in headings:
        return f"SECTION: {stripped}"
    return stripped


def _is_text_corrupted(text: str) -> bool:
    """
    Detects garbled text from symbol-encoded fonts (common with Greek letters
    like alpha/beta in statistics/math PDFs exported from PowerPoint). These
    show up as the Unicode replacement character or private-use-area codepoints
    that render as blank boxes.

    Triggers on ANY occurrence, not a percentage threshold — a single
    corrupted symbol in an otherwise clean slide is enough to justify
    re-extracting that page via OCR, since sparse corruption is the common
    real-world case for math/stats slides, not whole-page garbling.
    """
    if not text:
        return False

    for ch in text:
        codepoint = ord(ch)
        if ch == "\ufffd":
            return True
        if 0xE000 <= codepoint <= 0xF8FF:
            return True
        # Control characters (excluding normal whitespace) are another common
        # symptom of a broken font's character mapping outputting garbage.
        if codepoint < 32 and ch not in ("\n", "\t", "\r"):
            return True

    return False


def _extract_pdf_page_text_ordered(fitz_page, exclude_bboxes: list[tuple] = None, headings: set = None) -> str:
    exclude_bboxes = exclude_bboxes or []
    headings = headings or set()
    blocks = fitz_page.get_text("blocks")
    text_blocks = [b for b in blocks if b[4].strip() and b[6] == 0]

    def block_center(b):
        return ((b[0] + b[2]) / 2, (b[1] + b[3]) / 2)

    def is_excluded(b):
        cx, cy = block_center(b)
        return any(_point_in_bbox(cx, cy, bbox) for bbox in exclude_bboxes)

    text_blocks = [b for b in text_blocks if not is_excluded(b)]
    if not text_blocks:
        return ""

    page_width = fitz_page.rect.width
    full_width_threshold = 0.6 * page_width
    column_split_x = page_width / 2

    def block_width(b):
        return b[2] - b[0]

    narrow_blocks = [b for b in text_blocks if block_width(b) <= full_width_threshold]
    left_blocks = [b for b in narrow_blocks if (b[0] + b[2]) / 2 < column_split_x]
    right_blocks = [b for b in narrow_blocks if (b[0] + b[2]) / 2 >= column_split_x]

    is_multi_column = len(left_blocks) >= 2 and len(right_blocks) >= 2

    if not is_multi_column:
        ordered = sorted(text_blocks, key=lambda b: (b[1], b[0]))
        return "\n".join(_tag_if_heading(b[4], headings) for b in ordered)

    all_blocks_sorted_by_y = sorted(text_blocks, key=lambda b: b[1])
    output_lines = []
    pending_left, pending_right = [], []

    def flush_columns():
        pending_left.sort(key=lambda b: b[1])
        pending_right.sort(key=lambda b: b[1])
        for b in pending_left:
            output_lines.append(_tag_if_heading(b[4], headings))
        for b in pending_right:
            output_lines.append(_tag_if_heading(b[4], headings))
        pending_left.clear()
        pending_right.clear()

    for b in all_blocks_sorted_by_y:
        if block_width(b) > full_width_threshold:
            flush_columns()
            output_lines.append(_tag_if_heading(b[4], headings))
        else:
            x_center = (b[0] + b[2]) / 2
            if x_center < column_split_x:
                pending_left.append(b)
            else:
                pending_right.append(b)

    flush_columns()
    return "\n".join(line for line in output_lines if line)


def _ocr_page(fitz_page) -> str:
    try:
        zoom = 2
        matrix = fitz.Matrix(zoom, zoom)
        pixmap = fitz_page.get_pixmap(matrix=matrix)
        image_bytes = pixmap.tobytes("png")
        image = Image.open(io.BytesIO(image_bytes))
        text = pytesseract.image_to_string(image)
        return text.strip()
    except pytesseract.TesseractNotFoundError:
        raise ValueError(
            "OCR failed: Tesseract is not installed on this system. "
            "Install it with 'brew install tesseract' (Mac) or your OS's equivalent."
        )
    except Exception as e:
        raise ValueError(f"OCR processing failed: {e}")


class DocumentParser:
    SUPPORTED_EXTENSIONS = {
        ".pdf", ".docx", ".txt", ".md",
        ".xlsx", ".csv", ".json", ".html", ".htm", ".pptx",
    }

    def parse(self, filename: str, file_bytes: bytes) -> str:
        extension = self._get_extension(filename)

        if extension == ".pdf":
            return self._parse_pdf(file_bytes)
        elif extension == ".docx":
            return self._parse_docx(file_bytes)
        elif extension in (".txt", ".md"):
            return self._parse_plain_text(file_bytes)
        elif extension == ".xlsx":
            return self._parse_xlsx(file_bytes)
        elif extension == ".csv":
            return self._parse_csv(file_bytes)
        elif extension == ".json":
            return self._parse_json(file_bytes)
        elif extension in (".html", ".htm"):
            return self._parse_html(file_bytes)
        elif extension == ".pptx":
            return self._parse_pptx(file_bytes)
        else:
            raise ValueError(
                f"Unsupported file type: {extension}. Supported: {self.SUPPORTED_EXTENSIONS}"
            )

    def _parse_pdf(self, file_bytes: bytes) -> str:
        text_parts = []

        with fitz.open(stream=file_bytes, filetype="pdf") as fitz_doc, \
             pdfplumber.open(io.BytesIO(file_bytes)) as plumber_doc:

            for page_num, fitz_page in enumerate(fitz_doc):
                plumber_page = plumber_doc.pages[page_num]

                table_texts, table_bboxes = _extract_pdf_tables(plumber_page)
                headings = _detect_pdf_headings(fitz_page)
                prose_text = _extract_pdf_page_text_ordered(
                    fitz_page, exclude_bboxes=table_bboxes, headings=headings
                ).strip()

                combined_text = prose_text + " " + " ".join(table_texts)
                combined_length = len(combined_text.strip())
                is_corrupted = _is_text_corrupted(combined_text)

                if combined_length < OCR_MIN_TEXT_LENGTH or is_corrupted:
                    ocr_text = _ocr_page(fitz_page)
                    if ocr_text:
                        text_parts.append(f"[OCR] {ocr_text}")
                    elif prose_text or table_texts:
                        if prose_text:
                            text_parts.append(prose_text)
                        for table_text in table_texts:
                            text_parts.append(table_text)
                else:
                    if prose_text:
                        text_parts.append(prose_text)
                    for table_text in table_texts:
                        text_parts.append(table_text)

            embedded_images = _extract_pdf_images(fitz_doc)
            text_parts.extend(_describe_images_as_text(embedded_images))

        full_text = "\n\n".join(text_parts)
        if not full_text.strip():
            raise ValueError(
                "No extractable text found in PDF, even after attempting OCR."
            )
        return full_text

    def _parse_docx(self, file_bytes: bytes) -> str:
        doc = docx.Document(io.BytesIO(file_bytes))
        text_parts = []

        for block in _iter_block_items(doc):
            if isinstance(block, Paragraph):
                if block.text.strip():
                    if _is_docx_heading(block):
                        text_parts.append(f"SECTION: {block.text.strip()}")
                    else:
                        text_parts.append(block.text.strip())
            elif isinstance(block, Table):
                text_parts.append(_format_docx_table_as_text(block))

        embedded_images = _extract_docx_images(doc)
        text_parts.extend(_describe_images_as_text(embedded_images))

        full_text = "\n\n".join(text_parts)
        if not full_text.strip():
            raise ValueError("No extractable text found in DOCX file.")
        return full_text

    def _parse_plain_text(self, file_bytes: bytes) -> str:
        try:
            return file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            return file_bytes.decode("latin-1")

    def _parse_xlsx(self, file_bytes: bytes) -> str:
        workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        text_parts = []
        for sheet in workbook.worksheets:
            text_parts.append(f"Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                row_values = [str(cell) for cell in row if cell is not None]
                if row_values:
                    text_parts.append(" | ".join(row_values))
        full_text = "\n".join(text_parts)
        if not full_text.strip():
            raise ValueError("No extractable content found in XLSX file.")
        return full_text

    def _parse_csv(self, file_bytes: bytes) -> str:
        try:
            decoded = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            decoded = file_bytes.decode("latin-1")
        reader = csv.reader(io.StringIO(decoded))
        rows = [" | ".join(row) for row in reader if any(cell.strip() for cell in row)]
        full_text = "\n".join(rows)
        if not full_text.strip():
            raise ValueError("No extractable content found in CSV file.")
        return full_text

    def _parse_json(self, file_bytes: bytes) -> str:
        try:
            data = json.loads(file_bytes.decode("utf-8"))
        except json.JSONDecodeError as e:
            raise ValueError(f"Invalid JSON file: {e}")

        def flatten(obj, prefix=""):
            lines = []
            if isinstance(obj, dict):
                for key, value in obj.items():
                    lines.extend(flatten(value, f"{prefix}{key}."))
            elif isinstance(obj, list):
                for i, item in enumerate(obj):
                    lines.extend(flatten(item, f"{prefix}{i}."))
            else:
                lines.append(f"{prefix.rstrip('.')}: {obj}")
            return lines

        full_text = "\n".join(flatten(data))
        if not full_text.strip():
            raise ValueError("No extractable content found in JSON file.")
        return full_text

    def _parse_html(self, file_bytes: bytes) -> str:
        try:
            decoded = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            decoded = file_bytes.decode("latin-1")
        soup = BeautifulSoup(decoded, "html.parser")
        for tag in soup(["script", "style"]):
            tag.decompose()
        full_text = soup.get_text(separator="\n", strip=True)
        if not full_text.strip():
            raise ValueError("No extractable text found in HTML file.")
        return full_text

    def _parse_pptx(self, file_bytes: bytes) -> str:
        presentation = Presentation(io.BytesIO(file_bytes))
        text_parts = []
        for i, slide in enumerate(presentation.slides, start=1):
            slide_text = []
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    slide_text.append(shape.text_frame.text.strip())
            if slide_text:
                text_parts.append(f"Slide {i}: " + " | ".join(slide_text))
        embedded_images = _extract_pptx_images(presentation)
        text_parts.extend(_describe_images_as_text(embedded_images))
        full_text = "\n\n".join(text_parts)
        if not full_text.strip():
            raise ValueError("No extractable text found in PPTX file.")
        return full_text

    @staticmethod
    def _get_extension(filename: str) -> str:
        return "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


document_parser = DocumentParser()
